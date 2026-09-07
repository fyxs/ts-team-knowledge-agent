from __future__ import annotations
import re
from pathlib import Path
from openpyxl import load_workbook
DEFAULT_SHEET_ROWS=5000
def _clean(value: object) -> str:
    if value is None: return ""
    return re.sub(r"\s+", " ", str(value).replace("|", "\\|").replace("\r", " ").replace("\n", " ")).strip()
def convert_excel(source: Path, output: Path, max_rows: int=DEFAULT_SHEET_ROWS) -> int:
    if max_rows < 1: raise ValueError("Excel sheet row limit must be at least 1")
    output.parent.mkdir(parents=True, exist_ok=True); sheets_dir=output.parent/"sheets"; sheets_dir.mkdir(exist_ok=True)
    workbook=load_workbook(source, read_only=True, data_only=False); index=[f"# {source.stem}","","## Sheets",""]; total=0
    for ws in workbook.worksheets:
        rows=ws.iter_rows(values_only=True); header=next(rows,None)
        if header is None: continue
        headers=[_clean(v) or f"Column{i+1}" for i,v in enumerate(header)]; chunk=[]; part=1; files=[]; data_rows=0
        def flush():
            nonlocal chunk,part,total
            if not chunk: return
            safe=re.sub(r"[^0-9A-Za-z_-]+","_",ws.title).strip("_") or f"Sheet{ws._id}"; name=f"{safe}.md" if part==1 else f"{safe}-{part:03d}.md"; target=sheets_dir/name
            lines=[f"# {ws.title}" if part==1 else f"# {ws.title} (part {part})","","| "+" | ".join(headers)+" |","|"+"|".join("---" for _ in headers)+"|"]
            lines += ["| "+" | ".join(_clean(v) for v in row[:len(headers)])+" |" for row in chunk]; target.write_text("\n".join(lines)+"\n",encoding="utf-8"); files.append(target); total+=target.stat().st_size; chunk=[]; part+=1
        for row in rows:
            chunk.append(row); data_rows+=1
            if len(chunk)>=max_rows: flush()
        flush()
        if files: index.append(f"- [{ws.title}](sheets/{files[0].name}): {data_rows} rows, {len(files)} files")
    output.write_text("\n".join(index)+"\n",encoding="utf-8"); return output.stat().st_size+total
